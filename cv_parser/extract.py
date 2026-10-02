"""Turns an input file into Claude message content blocks (text and/or images)."""
import base64
from pathlib import Path

import pymupdf
from docx import Document

IMAGE_TYPES = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png"}
SUPPORTED = {".pdf", ".docx", *IMAGE_TYPES}
MIN_PDF_TEXT_CHARS = 200  # below this, treat the PDF as scanned and use vision
MAX_PDF_PAGES = 10


def _image_block(data: bytes, media_type: str) -> dict:
    return {
        "type": "image",
        "source": {"type": "base64", "media_type": media_type, "data": base64.b64encode(data).decode()},
    }


def _docx_text(path: Path) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip())
            if cells:
                parts.append(" | ".join(cells))
    for section in doc.sections:  # contact info often lives in the header
        parts.extend(p.text for p in section.header.paragraphs if p.text.strip())
    return "\n".join(parts)


def _pdf_blocks(path: Path) -> list[dict]:
    with pymupdf.open(path) as pdf:
        pages = list(pdf)[:MAX_PDF_PAGES]
        text = "\n".join(p.get_text() for p in pages)
        if len(text.strip()) >= MIN_PDF_TEXT_CHARS:
            return [{"type": "text", "text": text}]
        return [_image_block(p.get_pixmap(dpi=150).tobytes("png"), "image/png") for p in pages]


def to_content_blocks(path: str | Path) -> list[dict]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type '{ext}'. Use: {', '.join(sorted(SUPPORTED))}")
    if ext == ".pdf":
        return _pdf_blocks(path)
    if ext == ".docx":
        return [{"type": "text", "text": _docx_text(path)}]
    return [_image_block(path.read_bytes(), IMAGE_TYPES[ext])]
